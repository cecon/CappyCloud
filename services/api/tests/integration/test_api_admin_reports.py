"""Integração HTTP - /api/admin/reports (relatório de uso por workspace)."""

from __future__ import annotations

import csv
import io
import uuid
from collections.abc import AsyncGenerator
from datetime import UTC, datetime, timedelta

import pytest
from app.adapters.primary.http.admin_reports import get_workspace_report_repo
from app.adapters.primary.http.deps import get_db_session
from app.domain.entities import UserRole
from app.infrastructure.orm_models import Base, Conversation, Message, Sandbox
from app.infrastructure.orm_models import User as UserORM
from app.infrastructure.orm_models_workspaces import Workspace
from app.main import app as fastapi_app
from app.ports.workspace_report import ReportConversation, ReportMessage, ReportWorkspace
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from tests.conftest import InMemoryUserRepository
from tests.fakes_reports import InMemoryWorkspaceReportRepository
from tests.integration.conftest import seed_user

PROTEUS = ReportWorkspace(uuid.uuid4(), "proteus", "PROTEUS", "nfse-protheus")
SELLER = ReportWorkspace(uuid.uuid4(), "seller", "SELLER")
# 2026-09-01 12:00 em São Paulo; o período dos testes é 2026-09-01 a 2026-09-14.
T0 = datetime(2026, 9, 1, 15, tzinfo=UTC)
PERIOD = {"start": "2026-09-01", "end": "2026-09-14"}


def _conv(ws: ReportWorkspace, email: str, title: str, branch: str) -> ReportConversation:
    return ReportConversation(uuid.uuid4(), ws.id, email, title, "123", frozenset({branch}))


@pytest.fixture
def report_repo() -> AsyncGenerator[InMemoryWorkspaceReportRepository]:
    repo = InMemoryWorkspaceReportRepository()
    repo.add_workspace(PROTEUS)
    repo.add_workspace(SELLER)
    rows = [
        (_conv(PROTEUS, "ana@linx.com", "Rejeição E370 retenções", "main"), [(0, 1.0), (8, 0.5)]),
        (_conv(PROTEUS, "bruno@linx.com", "Falha de schema indDest", "main"), [(2, 2.0)]),
        (_conv(PROTEUS, "ana@linx.com", "ola", "release"), [(-2, 7.0), (3, 0.25)]),
        (_conv(SELLER, "carla@linx.com", "como ativar o sap", "master"), [(1, 4.0)]),
    ]
    for conv, msgs in rows:
        repo.add_conversation(
            conv, *(ReportMessage(conv.id, T0 + timedelta(days=d), c) for d, c in msgs)
        )
    fastapi_app.dependency_overrides[get_workspace_report_repo] = lambda: repo
    yield repo
    fastapi_app.dependency_overrides.pop(get_workspace_report_repo, None)


async def _headers(
    client: AsyncClient, user_repo: InMemoryUserRepository, *, super_admin: bool = False
) -> tuple[dict[str, str], uuid.UUID]:
    email = f"admin-{uuid.uuid4().hex[:6]}@test.com"
    user = await seed_user(user_repo, email, role=UserRole.ADMIN, is_super_admin=super_admin)
    login = await client.post(
        "/api/auth/login", data={"username": email, "password": "password123"}
    )
    return {"Authorization": f"Bearer {login.json()['access_token']}"}, user.id


async def test_report_metrics_for_one_workspace(
    client: AsyncClient,
    user_repo: InMemoryUserRepository,
    report_repo: InMemoryWorkspaceReportRepository,
) -> None:
    headers, _ = await _headers(client, user_repo, super_admin=True)
    res = await client.get(
        "/api/admin/reports/workspace",
        headers=headers,
        params={"workspace_id": str(PROTEUS.id), **PERIOD},
    )
    assert res.status_code == 200
    body = res.json()
    assert body["filters"]["workspace_name"] == "PROTEUS"
    assert body["filters"]["currency"] == "USD"
    assert body["filters"]["timezone"] == "America/Sao_Paulo"
    assert body["totals"] == {
        "consultations": 3,
        "analysts": 2,
        "messages": 4,
        "cost_usd": 3.75,
        "avg_cost_per_consultation": 1.25,
        "avg_cost_per_analyst": 1.875,
    }
    assert [(a["label"], a["consultations"], a["cost_usd"]) for a in body["analysts"]] == [
        ("bruno", 1, 2.0),
        ("ana", 2, 1.75),
    ]
    assert [(w["week"], w["consultations"], w["cost_usd"]) for w in body["weeks"]] == [
        ("2026-W36", 3, 3.25),
        ("2026-W37", 1, 0.5),
        ("2026-W38", 0, 0.0),
    ]
    assert [(t["key"], t["consultations"]) for t in body["themes"]] == [
        ("nfse-schema-xml", 1),
        ("nfse-retencoes", 1),
        ("outros", 1),
    ]
    assert len(body["consultations"]) == 3


async def test_branch_and_period_filters(
    client: AsyncClient,
    user_repo: InMemoryUserRepository,
    report_repo: InMemoryWorkspaceReportRepository,
) -> None:
    headers, _ = await _headers(client, user_repo, super_admin=True)
    params = {"workspace_id": str(PROTEUS.id), "branch": "release", **PERIOD}
    body = (await client.get("/api/admin/reports/workspace", headers=headers, params=params)).json()
    assert [c["title"] for c in body["consultations"]] == ["ola"]
    # A mensagem de 7.0 é anterior ao período: só 0.25 conta.
    assert body["totals"]["cost_usd"] == 0.25
    params = {"start": "2026-09-02", "end": "2026-09-02"}
    body = (await client.get("/api/admin/reports/workspace", headers=headers, params=params)).json()
    assert [c["workspace_slug"] for c in body["consultations"]] == ["seller"]


async def test_invalid_period_returns_422(
    client: AsyncClient,
    user_repo: InMemoryUserRepository,
    report_repo: InMemoryWorkspaceReportRepository,
) -> None:
    headers, _ = await _headers(client, user_repo, super_admin=True)
    for params in (
        {"start": "2026-09-10", "end": "2026-09-01"},
        {"start": "2025-01-01", "end": "2026-09-01"},
    ):
        res = await client.get("/api/admin/reports/workspace", headers=headers, params=params)
        assert res.status_code == 422


async def test_regular_admin_is_restricted_to_granted_workspaces(
    client: AsyncClient,
    user_repo: InMemoryUserRepository,
    report_repo: InMemoryWorkspaceReportRepository,
) -> None:
    headers, user_id = await _headers(client, user_repo)
    report_repo.grant(user_id, PROTEUS.id)
    seller = {"workspace_id": str(SELLER.id), **PERIOD}
    for path in ("/workspace", "/workspace/export", "/options"):
        res = await client.get(f"/api/admin/reports{path}", headers=headers, params=seller)
        assert res.status_code == 403, path
    res = await client.get(f"/api/admin/reports/themes/{SELLER.id}", headers=headers)
    assert res.status_code == 403
    body = (await client.get("/api/admin/reports/workspace", headers=headers, params=PERIOD)).json()
    assert [w["slug"] for w in body["workspaces"]] == ["proteus"]
    assert body["totals"]["consultations"] == 3
    options = (await client.get("/api/admin/reports/options", headers=headers)).json()
    assert [w["slug"] for w in options["workspaces"]] == ["proteus"]
    assert options["branches"] == ["main", "release"]
    assert options["can_edit_themes"] is False
    res = await client.get(
        "/api/admin/reports/workspace", headers=headers, params={"workspace_id": str(uuid.uuid4())}
    )
    assert res.status_code == 404


async def test_non_admin_gets_403(
    client: AsyncClient,
    user_headers: dict[str, str],
    report_repo: InMemoryWorkspaceReportRepository,
) -> None:
    for path in ("/workspace", "/options", f"/themes/{PROTEUS.id}"):
        res = await client.get(f"/api/admin/reports{path}", headers=user_headers)
        assert res.status_code == 403


async def test_theme_editing_is_super_admin_only(
    client: AsyncClient,
    user_repo: InMemoryUserRepository,
    report_repo: InMemoryWorkspaceReportRepository,
) -> None:
    admin, admin_id = await _headers(client, user_repo)
    report_repo.grant(admin_id, PROTEUS.id)
    url = f"/api/admin/reports/themes/{PROTEUS.id}"
    rules = {"rules": [{"key": "saudacao", "label": "Saudação", "patterns": ["\\bola\\b"]}]}
    assert (await client.put(url, headers=admin, json=rules)).status_code == 403
    root, _ = await _headers(client, user_repo, super_admin=True)
    current = (await client.get(url, headers=root)).json()
    assert (current["preset"], current["custom"]) == ("nfse-protheus", False)
    assert [p["key"] for p in current["presets"]] == ["generico", "nfse-protheus"]
    bad = await client.put(
        url, headers=root, json={"rules": [{"key": "x", "label": "X", "patterns": ["("]}]}
    )
    assert bad.status_code == 422
    assert "padrão inválido" in bad.json()["detail"]
    saved = (await client.put(url, headers=root, json=rules)).json()
    assert (saved["custom"], saved["rules"][0]["key"]) == (True, "saudacao")
    body = (
        await client.get(
            "/api/admin/reports/workspace",
            headers=admin,
            params={"workspace_id": str(PROTEUS.id), **PERIOD},
        )
    ).json()
    assert [(t["key"], t["consultations"]) for t in body["themes"]] == [
        ("saudacao", 1),
        ("outros", 2),
    ]
    assert (await client.put(url, headers=root, json={"preset": "x"})).status_code == 422


async def test_exports_match_report(
    client: AsyncClient,
    user_repo: InMemoryUserRepository,
    report_repo: InMemoryWorkspaceReportRepository,
) -> None:
    headers, _ = await _headers(client, user_repo, super_admin=True)
    params = {"workspace_id": str(PROTEUS.id), **PERIOD}
    res = await client.get(
        "/api/admin/reports/workspace/export", headers=headers, params={**params, "format": "csv"}
    )
    assert res.status_code == 200
    assert (
        'filename="relatorio-proteus-2026-09-01-2026-09-14.csv"'
        in res.headers["content-disposition"]
    )
    rows = list(csv.reader(io.StringIO(res.content.decode("utf-8-sig")), delimiter=";"))
    assert rows[0][:3] == ["workspace", "analista", "chamado"]
    assert sorted(float(r[7]) for r in rows[1:]) == [0.25, 1.5, 2.0]

    res = await client.get("/api/admin/reports/workspace/export", headers=headers, params=PERIOD)
    assert res.status_code == 200
    assert "relatorio-todos-" in res.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["Resumo", "Analistas", "Semanas", "Temas", "Workspaces", "Consultas"]
    summary = {row[0]: row[1] for row in wb["Resumo"].iter_rows(min_row=2, values_only=True)}
    assert summary["Consultas"] == 4
    assert summary["Custo total (US$)"] == 7.75
    assert "messages.cost_usd" in summary["Fonte do custo"]
    assert wb["Consultas"].max_row == 5


@pytest.fixture
async def sql_session() -> AsyncGenerator[AsyncSession]:
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    async with async_sessionmaker(engine, expire_on_commit=False)() as session:
        fastapi_app.dependency_overrides[get_db_session] = lambda: session
        yield session
        fastapi_app.dependency_overrides.pop(get_db_session, None)
    await engine.dispose()


async def test_real_repository_wiring(
    client: AsyncClient, user_repo: InMemoryUserRepository, sql_session: AsyncSession
) -> None:
    headers, _ = await _headers(client, user_repo, super_admin=True)
    sandbox = Sandbox(id=uuid.uuid4(), name="sb", host="sb", session_port=8080)
    user = UserORM(id=uuid.uuid4(), email="ana@linx.com", hashed_password="x")
    ws = Workspace(
        id=uuid.uuid4(),
        slug="proteus",
        name="PROTEUS",
        sandbox_id=sandbox.id,
        report_theme_preset="nfse-protheus",
    )
    conv = Conversation(
        id=uuid.uuid4(),
        user_id=user.id,
        workspace_id=ws.id,
        title="Falha de schema",
        repos=[{"base_branch": "main"}],
    )
    sql_session.add_all([sandbox, user])
    await sql_session.flush()
    sql_session.add(ws)
    await sql_session.flush()
    sql_session.add(conv)
    await sql_session.flush()
    sql_session.add(
        Message(
            id=uuid.uuid4(),
            conversation_id=conv.id,
            role="assistant",
            content="ok",
            cost_usd=0.5,
            created_at=T0,
        )
    )
    await sql_session.commit()

    body = (await client.get("/api/admin/reports/workspace", headers=headers, params=PERIOD)).json()
    assert body["totals"]["cost_usd"] == 0.5
    assert body["consultations"][0]["theme_key"] == "nfse-schema-xml"
    assert body["consultations"][0]["branches"] == ["main"]
    options = (await client.get("/api/admin/reports/options", headers=headers)).json()
    assert options["branches"] == ["main"]
