"""Arquivos do relatório de uso: XLSX (abas para a Diretoria) e CSV (consultas)."""

from __future__ import annotations

import codecs
import csv
import io
from collections.abc import Sequence
from datetime import datetime
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from app.application.use_cases.workspace_report import BuiltReport

USD_FORMAT = '"US$" #,##0.0000'
COST_SOURCE = "Valores em US$, somados de messages.cost_usd (custo real informado pelo provedor)."

CONSULTATION_COLUMNS = [
    "workspace",
    "analista",
    "chamado",
    "titulo",
    "tema",
    "branches",
    "mensagens_no_periodo",
    "custo_usd",
    "primeira_mensagem",
    "ultima_mensagem",
]


def export_filename(built: BuiltReport, ext: str) -> str:
    return f"relatorio-{built.scope_slug}-{built.start.isoformat()}-{built.end.isoformat()}.{ext}"


def _local(value: datetime) -> datetime:
    """openpyxl não grava fuso; a data já está no fuso do relatório."""
    return value.replace(tzinfo=None)


def _consultation_values(built: BuiltReport) -> list[list[Any]]:
    return [
        [
            row.workspace_slug,
            row.analyst_email,
            row.ticket_number or "",
            row.title,
            row.theme_label,
            ", ".join(row.branches),
            row.messages,
            row.cost_usd,
            _local(row.first_message_at),
            _local(row.last_message_at),
        ]
        for row in built.report.consultations
    ]


def report_csv(built: BuiltReport) -> bytes:
    buffer = io.StringIO()
    # Ponto e vírgula: o Excel em português abre direto em colunas.
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(CONSULTATION_COLUMNS)
    for values in _consultation_values(built):
        values[7] = f"{values[7]:.6f}"
        values[8] = values[8].isoformat(timespec="seconds")
        values[9] = values[9].isoformat(timespec="seconds")
        writer.writerow(values)
    # BOM para o Excel reconhecer UTF-8 (acentos dos títulos).
    return codecs.BOM_UTF8 + buffer.getvalue().encode("utf-8")


def _sheet(
    wb: Workbook,
    title: str,
    header: Sequence[str],
    rows: Sequence[Sequence[Any]],
    money_cols: Sequence[int] = (),
) -> Worksheet:
    ws = wb.create_sheet(title)
    ws.append(list(header))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for values in rows:
        ws.append(list(values))
    for col in money_cols:
        for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cell.number_format = USD_FORMAT
    for idx, name in enumerate(header, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = max(
            14, min(60, len(name) + 4)
        )
    return ws


def _summary_rows(built: BuiltReport) -> list[list[Any]]:
    totals = built.report.totals
    return [
        ["Workspace", built.workspace_name or "Todos os workspaces visíveis"],
        ["Branch", built.branch or "Todas"],
        ["Período", f"{built.start.isoformat()} a {built.end.isoformat()} ({built.timezone})"],
        ["Gerado em", _local(built.generated_at.astimezone(ZoneInfo(built.timezone)))],
        ["Consultas", totals.consultations],
        ["Analistas distintos", totals.analysts],
        ["Mensagens no período", totals.messages],
        ["Custo total (US$)", totals.cost_usd],
        ["Custo médio por consulta (US$)", totals.avg_cost_per_consultation],
        ["Custo médio por analista (US$)", totals.avg_cost_per_analyst],
        ["Fonte do custo", COST_SOURCE],
    ]


def report_xlsx(built: BuiltReport) -> bytes:
    report = built.report
    wb = Workbook()
    wb.remove(wb.active)
    summary = _sheet(wb, "Resumo", ["Indicador", "Valor"], _summary_rows(built))
    for row in (9, 10, 11):
        summary.cell(row=row, column=2).number_format = USD_FORMAT
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 60
    _sheet(
        wb,
        "Analistas",
        ["Analista", "E-mail", "Consultas", "Custo (US$)", "Custo médio (US$)"],
        [[a.label, a.email, a.consultations, a.cost_usd, a.avg_cost_usd] for a in report.analysts],
        money_cols=(4, 5),
    )
    _sheet(
        wb,
        "Semanas",
        ["Semana ISO", "Início (segunda)", "Consultas", "Custo (US$)"],
        [[w.week, w.start, w.consultations, w.cost_usd] for w in report.weeks],
        money_cols=(4,),
    )
    _sheet(
        wb,
        "Temas",
        ["Tema", "Consultas", "Custo (US$)"],
        [[t.label, t.consultations, t.cost_usd] for t in report.themes],
        money_cols=(3,),
    )
    if built.workspace_id is None:
        _sheet(
            wb,
            "Workspaces",
            ["Workspace", "Consultas", "Custo (US$)"],
            [[w.name, w.consultations, w.cost_usd] for w in report.workspaces],
            money_cols=(3,),
        )
    _sheet(wb, "Consultas", CONSULTATION_COLUMNS, _consultation_values(built), money_cols=(8,))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
